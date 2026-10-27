#!/usr/bin/env python3
"""SwimSwam Daily Swim Coach Workout index builder.

Collects links + metadata (Purpose, Target age group, Target level, Weeks until
target meet, Course) for each workout post and writes a formatted Excel file.
It does NOT save workout text.

Politeness: honest User-Agent, obeys robots.txt, rate-limited, and it stops
(rather than evading) if the site blocks it. See README for the browser-saved
pages fallback (LOCAL_DIR).
"""
import csv
import os
import random
import re
import sys
import time
import urllib.robotparser
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---- Configuration (environment variables) --------------------------------
START = int(os.environ.get("START") or 100)
END = int(os.environ.get("END") or 1136)
DELAY_MIN = float(os.environ.get("DELAY_MIN") or 4)
DELAY_MAX = float(os.environ.get("DELAY_MAX") or 8)
COURSE_FILTER = os.environ.get("COURSE_FILTER", "25 Yards")  # blank = keep all
OUTPUT = os.environ.get("OUTPUT") or "swimswam_workouts.xlsx"
LOCAL_DIR = os.environ.get("LOCAL_DIR")  # folder of HTML files saved from a browser
CONTACT = os.environ.get("CONTACT") or "unset"
MAX_BLOCKS = int(os.environ.get("MAX_BLOCKS") or 3)

BASE_URL = "https://swimswam.com/daily-swim-coach-workout-{n}/"
USER_AGENT = f"swimswam-workout-index/0.1 (personal index builder; contact: {CONTACT})"

LABELS = {
    "purpose": ["purpose"],
    "age_group": ["target age group", "age group"],
    "level": ["target level"],
    "weeks": ["weeks until target meet", "weeks until meet"],
    "course": ["course"],
}
_ALL_LABELS = {l for ls in LABELS.values() for l in ls}


# ---- Parsing ---------------------------------------------------------------
def _lines(html):
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        t.decompose()
    root = soup.find("article") or soup.body or soup
    lines = [re.sub(r"\s+", " ", l).strip() for l in root.get_text("\n").split("\n")]
    return soup, [l for l in lines if l]


def parse_page(html):
    """Return dict with title + the five metadata fields (None if not found)."""
    soup, lines = _lines(html)
    h1 = soup.find("h1")
    out = {"title": h1.get_text(strip=True) if h1 else (soup.title.get_text(strip=True) if soup.title else None)}
    for field, names in LABELS.items():
        out[field] = None
        inline = re.compile(r"^(?:%s)\s*[:\-\u2013]\s*(.+)$" % "|".join(map(re.escape, names)), re.I)
        for i, line in enumerate(lines):
            m = inline.match(line)
            if m:
                out[field] = m.group(1).strip()
                break
            if line.lower().rstrip(":").strip() in names and i + 1 < len(lines):
                nxt = lines[i + 1]
                if nxt.lower().rstrip(":").strip() not in _ALL_LABELS:
                    out[field] = nxt.strip()
                    break
    if out["weeks"]:
        m = re.search(r"\d+(?:\s*[-\u2013]\s*\d+)?", out["weeks"])
        out["weeks"] = m.group(0).replace(" ", "") if m else out["weeks"]
    return out


def has_metadata(d):
    return any(d.get(k) for k in LABELS)


# ---- Collecting ------------------------------------------------------------
def collect_local(folder):
    rows = []
    for p in sorted(Path(folder).glob("*.htm*")):
        m = re.search(r"(\d+)(?!.*\d)", p.stem)
        if not m:
            print(f"skip {p.name}: no workout number in filename")
            continue
        n = int(m.group(1))
        d = parse_page(p.read_text(encoding="utf-8", errors="ignore"))
        rows.append({"number": n, "url": BASE_URL.format(n=n), "status": "ok" if has_metadata(d) else "no metadata found", **d})
        print(f"#{n}: {rows[-1]['status']}")
    return rows


def collect_web():
    s = requests.Session()
    s.headers["User-Agent"] = USER_AGENT
    # Fetch robots.txt ourselves so a block is reported as a block, not as
    # "everything disallowed" (urllib's parser treats a 403 that way).
    rp = urllib.robotparser.RobotFileParser()
    try:
        rr = s.get("https://swimswam.com/robots.txt", timeout=30)
    except requests.RequestException as e:
        rr = None
        print(f"Could not fetch robots.txt ({e.__class__.__name__}); continuing cautiously")
    if rr is not None and rr.status_code == 200:
        rp.parse(rr.text.splitlines())
    elif rr is not None and rr.status_code in (401, 403):
        print(f"robots.txt request was blocked (HTTP {rr.status_code}). "
              "The site is refusing automated requests from this machine; stopping.")
        return [{"number": m, "url": BASE_URL.format(n=m), "title": None,
                 **{k: None for k in LABELS},
                 "status": f"blocked (HTTP {rr.status_code} on robots.txt)"}
                for m in range(START, END + 1)]
    else:
        rp = None
    rows, blocks = [], 0
    for n in range(START, END + 1):
        url = BASE_URL.format(n=n)
        row = {"number": n, "url": url, "title": None, **{k: None for k in LABELS}, "status": ""}
        if rp and not rp.can_fetch(USER_AGENT, url):
            row["status"] = "disallowed by robots.txt"
            rows.append(row)
            continue
        try:
            r = s.get(url, timeout=30)
        except requests.RequestException as e:
            row["status"] = f"network error: {e.__class__.__name__}"
            rows.append(row)
            continue
        if r.status_code == 429:
            wait = int(r.headers.get("Retry-After", "60"))
            print(f"#{n}: 429, waiting {wait}s")
            time.sleep(wait)
            row["status"] = "rate limited"
        elif r.status_code in (401, 403):
            blocks += 1
            row["status"] = f"blocked (HTTP {r.status_code})"
            print(f"#{n}: HTTP {r.status_code}")
            if blocks >= MAX_BLOCKS:
                rows.append(row)
                for m in range(n + 1, END + 1):
                    rows.append({"number": m, "url": BASE_URL.format(n=m), "title": None,
                                 **{k: None for k in LABELS}, "status": "not attempted"})
                print("Site is blocking automated requests; stopping. "
                      "Use LOCAL_DIR with pages saved from your browser instead.")
                return rows
        elif r.status_code == 404:
            blocks = 0
            row["status"] = "not found"
        elif r.ok:
            blocks = 0
            row.update(parse_page(r.text))
            row["status"] = "ok" if has_metadata(row) else "no metadata found"
            print(f"#{n}: {row['status']}")
        else:
            row["status"] = f"HTTP {r.status_code}"
        rows.append(row)
        time.sleep(random.uniform(DELAY_MIN, DELAY_MAX))
    return rows


# ---- Excel output ----------------------------------------------------------
HEADERS = ["Workout #", "Title", "URL", "Purpose", "Target age group",
           "Target level", "Weeks until target meet", "Course"]
KEYS = ["number", "title", "url", "purpose", "age_group", "level", "weeks", "course"]
WIDTHS = [11, 40, 52, 28, 22, 20, 24, 14]


def _sheet(ws, rows, with_status=False):
    heads = HEADERS + (["Status"] if with_status else [])
    keys = KEYS + (["status"] if with_status else [])
    widths = WIDTHS + ([26] if with_status else [])
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(k) for k in keys])
        cell = ws.cell(row=ws.max_row, column=3)
        cell.hyperlink = r["url"]
        cell.font = Font(color="0563C1", underline="single")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_excel(rows, path):
    rows = sorted(rows, key=lambda r: r["number"])
    if COURSE_FILTER:
        kept = [r for r in rows if r.get("course") and COURSE_FILTER.lower() in r["course"].lower()]
    else:
        kept = rows
    wb = Workbook()
    ws = wb.active
    ws.title = "Workouts" + (f" ({COURSE_FILTER})" if COURSE_FILTER else "")
    ws.title = ws.title[:31]
    _sheet(ws, kept)
    _sheet(wb.create_sheet("All rows"), rows, with_status=True)
    wb.save(path)
    ok = sum(1 for r in rows if r["status"] == "ok")
    print(f"Saved {path}: {len(kept)} in main sheet, {len(rows)} total rows, {ok} parsed OK")


# ---- Repo index files (INDEX.md + data/workout_index.csv) -------------------
INDEX_HEADERS = ["Workout #", "Purpose", "Target age group", "Target level",
                 "Weeks until target meet", "Course"]


def write_index(rows, md_path="INDEX.md", csv_path="data/workout_index.csv"):
    """Write the repo's index: a Markdown table GitHub renders, and a CSV."""
    rows = sorted(rows, key=lambda r: r["number"])
    cell = lambda v: (v or "").replace("|", "/")
    lines = ["# SwimSwam Daily Swim Coach Workout Index", "",
             "Links and metadata only. Regenerate with `python scraper.py`.", "",
             "| " + " | ".join(INDEX_HEADERS) + " |",
             "|" + "|".join(["---"] * len(INDEX_HEADERS)) + "|"]
    for r in rows:
        link = f"[{r['number']}]({r['url']})"
        lines.append("| " + " | ".join([link, cell(r.get("purpose")), cell(r.get("age_group")),
                     cell(r.get("level")), cell(r.get("weeks")), cell(r.get("course"))]) + " |")
    Path(md_path).write_text("\n".join(lines) + "\n", encoding="utf-8")
    Path(csv_path).parent.mkdir(parents=True, exist_ok=True)
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Workout #", "URL", "Purpose", "Target age group", "Target level",
                    "Weeks until target meet", "Course"])
        for r in rows:
            w.writerow([r["number"], r["url"], r.get("purpose") or "", r.get("age_group") or "",
                        r.get("level") or "", r.get("weeks") or "", r.get("course") or ""])
    print(f"Wrote {md_path} and {csv_path} ({len(rows)} rows)")


def main():
    rows = collect_local(LOCAL_DIR) if LOCAL_DIR else collect_web()
    if not rows:
        sys.exit("No rows collected.")
    write_excel(rows, OUTPUT)
    write_index(rows)
    if any(str(r["status"]).startswith(("blocked", "disallowed")) for r in rows):
        sys.exit("FAILED: SwimSwam blocked or disallowed the requests (see log). "
                 "Use LOCAL_DIR with pages saved from your browser.")


if __name__ == "__main__":
    main()
